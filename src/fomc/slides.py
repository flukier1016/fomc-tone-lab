"""`fomc slides`: the Story as PowerPoint files the team can present, edit and restyle.

Written to --out (docs/slides by default):
  FOMC-Tone-Lab.pptx            every slide exactly as the Story page renders it (16:9 images), with real
                                slide titles for navigation and the speaker notes, ready to present;
  FOMC-Tone-Lab-editable.pptx   the same slides rebuilt from native text boxes (title, number, chips,
                                takeaway, source) with each chart or diagram as a picture, for restyling
                                in any template;
  outline.md                    every title, number, takeaway and speaker note as plain text;
  contact-sheet.jpg             every slide as a thumbnail on one image, the deck at a glance.

Rendering uses headless Chrome (or Chromium) against a copy of the dashboard started for the purpose,
so the files always match what the Story page shows for the current store.
"""

from __future__ import annotations

import json
import shutil
import socket
import subprocess
import tempfile
import threading
import time
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from datetime import UTC, datetime
from pathlib import Path

STAGE_W, STAGE_H = 1600, 900
INK, TEXT2, MUTED = "110F0E", "4B4743", "67625D"  # the dashboard's light tokens ("Ink and Vermilion")
ACCENT, ACCENT_SOFT, ACCENT_INK, SURFACE2, BORDER = "151312", "E8E4DC", "151312", "EEEBE5", "DDD9D0"
FONT = "Arial"  # present on every Mac and Windows machine, and the dashboard's own text face
SERIF = "Georgia"  # likewise; the dashboard sets its titles and headings in it too


# ----------------------------------------------------------------------- rendering
def chrome_path() -> str:
    for c in (
        "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
        "google-chrome",
        "chromium",
        "chromium-browser",
    ):
        if Path(c).exists() or shutil.which(c):
            return c
    raise SystemExit("Google Chrome or Chromium is needed to render the slides")


def shot(url: str, out: Path, size: tuple[int, int], scale: float = 2) -> Path:
    with tempfile.TemporaryDirectory() as profile:
        args = [chrome_path(), "--headless=new", "--disable-gpu", "--hide-scrollbars", f"--user-data-dir={profile}"]
        args += [f"--force-device-scale-factor={scale}", f"--window-size={size[0]},{size[1]}"]
        args += ["--virtual-time-budget=10000", f"--screenshot={out}", url]
        subprocess.run(args, check=True, capture_output=True, timeout=180)
    return out


class LocalServer:
    """The dashboard on a free local port, in a background thread, for as long as the `with` block runs."""

    def __enter__(self):
        import uvicorn

        sock = socket.socket()
        sock.bind(("127.0.0.1", 0))
        self.port = sock.getsockname()[1]
        sock.close()
        self.server = uvicorn.Server(
            uvicorn.Config("fomc.api:app", host="127.0.0.1", port=self.port, log_level="warning")
        )
        self.thread = threading.Thread(target=self.server.run, daemon=True)
        self.thread.start()
        for _ in range(200):
            try:
                urllib.request.urlopen(f"{self.base}/api/health", timeout=1)
                break
            except OSError:
                time.sleep(0.1)
        return self

    @property
    def base(self) -> str:
        return f"http://127.0.0.1:{self.port}"

    def __exit__(self, *exc):
        self.server.should_exit = True
        self.thread.join(timeout=10)


# ------------------------------------------------------------------------ builders
def _px(v: float):
    from pptx.util import Emu

    return Emu(int(v * 914400 * 13.333 / STAGE_W))  # stage pixels to EMU on a 13.333-inch slide


def _rgb(hex_: str):
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(hex_)


def _text(slide, x, y, w, h, text, size, color=INK, bold=False, align=None, caps=False, spacing=None, font=FONT):
    from pptx.enum.text import MSO_ANCHOR
    from pptx.util import Pt

    box = slide.shapes.add_textbox(_px(x), _px(y), _px(w), _px(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    if align is not None:
        p.alignment = align
    if spacing:
        p.line_spacing = spacing
    run = p.add_run()
    run.text = text.upper() if caps else text
    run.font.name, run.font.size, run.font.bold = font, Pt(size), bold
    run.font.color.rgb = _rgb(color)
    return box


def _pill(slide, x, y, w, h, text, size, fill=SURFACE2, color=INK, bold=False, radius=0.25):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR
    from pptx.util import Pt

    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, _px(x), _px(y), _px(w), _px(h))
    shp.adjustments[0] = radius
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(fill)
    shp.line.color.rgb = _rgb(BORDER)
    shp.line.width = Pt(0.75)
    tf = shp.text_frame
    tf.word_wrap = radius < 0.5  # name pills stay on one line; chips wrap
    tf.margin_left = tf.margin_right = _px(14)
    tf.margin_top = tf.margin_bottom = _px(6)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    run = tf.paragraphs[0].add_run()
    run.text = text
    run.font.name, run.font.size, run.font.bold = FONT, Pt(size), bold
    run.font.color.rgb = _rgb(color)
    return shp


def _notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s.get("notes", "")


def _nav_title(slide, title):
    """A real slide title for the outline, search and screen readers, placed just above the slide."""
    if slide.shapes.title is not None:
        t = slide.shapes.title
        t.text = title
        t.left, t.top, t.width, t.height = 0, -_px(140), _px(STAGE_W), _px(100)


def _properties(prs, story):
    """Use the project's own identity instead of the presentation library's template metadata."""
    props = prs.core_properties
    props.title = story["slides"][0]["title"]
    props.subject = "Federal Reserve communication, Treasury yields and duration signals"
    props.author = ", ".join(story["team"])
    props.last_modified_by = "Final Project Group 2"
    props.comments = "Research findings and technical implementation from the frozen FOMC Tone Lab store."
    props.created = props.modified = datetime.now(UTC).replace(tzinfo=None)


def image_deck(story: dict, images: dict[int, Path], out: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    _properties(prs, story)
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for s in story["slides"]:
        slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only: the title moves off the canvas
        _nav_title(slide, s["title"])
        pic = slide.shapes.add_picture(str(images[s["n"]]), 0, 0, prs.slide_width, prs.slide_height)
        pic.name = f"Slide {s['n']}"
        _notes(slide, s)
    prs.save(out)
    return out


def editable_deck(story: dict, visuals: dict[int, Path], out: Path, logo: Path | None = None) -> Path:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches

    prs = Presentation()
    _properties(prs, story)
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    total = len(story["slides"])
    for s in story["slides"]:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _rgb("FFFEFB")
        lay = s.get("layout", "standard")
        if lay in ("cover", "closing"):
            compact = s.get("id") == "cover"
            if logo:
                slide.shapes.add_picture(str(logo), _px(128), _px(220 if compact else 128), _px(84), _px(84))
            _text(slide, 128, 348 if compact else 250, 1340, 30, s["kicker"], 13, ACCENT, bold=True, caps=True)
            _text(
                slide,
                128,
                390 if compact else 300,
                1340,
                120 if compact else 220 if lay == "cover" else 160,
                s["title"],
                54 if lay == "cover" else 72,
                bold=True,
                font=SERIF,
            )
            _text(slide, 128, 520 if compact else 540 if lay == "cover" else 470, 1200, 90, s["subtitle"], 19, TEXT2)
            x = 128
            for name in s.get("team", []):
                w = 60 + 15 * len(name)
                _pill(
                    slide, x, 615 if compact else 660 if lay == "cover" else 570, w, 52, name, 14, bold=True, radius=0.5
                )
                x += w + 14
            if s.get("url"):
                _text(slide, 128, 650, 1200, 34, s["url"], 14, ACCENT_INK, bold=True)
        elif lay == "divider":
            _text(slide, 110, 250, 560, 380, s["number"], 170, ACCENT, bold=True, font=SERIF)
            _text(slide, 700, 270, 820, 30, s["kicker"], 13, ACCENT, bold=True, caps=True)
            _text(slide, 700, 310, 820, 100, s["title"], 44, bold=True, font=SERIF)
            _text(slide, 700, 420, 820, 50, s.get("subtitle", ""), 17, TEXT2)
        else:
            _text(slide, 72, 56, 1200, 28, s["kicker"], 11, ACCENT, bold=True, caps=True)
            _text(slide, 1328, 56, 200, 28, f"{s['n']} / {total}", 11, TEXT2, align=PP_ALIGN.RIGHT)
            _text(slide, 72, 96, 1420, 60, s["title"], 28, bold=True, spacing=1.0, font=SERIF)
            _text(slide, 72, 164, 1420, 56, s.get("subtitle", ""), 13.5, TEXT2)
            chips = [c for c in s.get("chips", []) if c]
            if lay in ("standard", "mirror"):  # the number and its facts in a column, left or right of the visual
                x = 72 if lay == "standard" else 1128
                _text(slide, x, 250, 400, 110, s.get("number", ""), 54, ACCENT, bold=True, font=SERIF)
                _text(slide, x, 362, 400, 100, s.get("number_label", ""), 13, TEXT2)
                y = 470
                for chip in chips[:4]:
                    _pill(slide, x, y, 400, 54, chip, 11, radius=0.2)
                    y += 62
            elif lay == "banner":  # the number beside its label, then the facts in a row
                _text(slide, 72, 250, 230, 100, s.get("number", ""), 44, ACCENT, bold=True, font=SERIF)
                _text(slide, 310, 252, 300, 100, s.get("number_label", ""), 11.5, TEXT2)
                row = chips[:4]
                w = (880 - 12 * (len(row) - 1)) / max(1, len(row))
                for k, chip in enumerate(row):
                    _pill(slide, 648 + k * (w + 12), 254, w, 96, chip, 10.5, radius=0.15)
            elif lay == "compare":  # two numbers set against each other
                accent = 1 if s.get("compare_word") == "but" else 0
                y = 262
                for k, (num, label) in enumerate(s.get("compare") or []):
                    if k:
                        word = s.get("compare_word", "against")
                        _text(slide, 72, y, 400, 28, word, 10.5, TEXT2, bold=True, caps=True)
                        y += 40
                    _text(slide, 72, y, 400, 90, num, 48, ACCENT if k == accent else TEXT2, bold=True, font=SERIF)
                    _text(slide, 72, y + 102, 400, 60, label, 13, TEXT2)
                    y += 180
            if s["n"] in visuals:
                boxes = {"wide": (72, 244, 1456, 500), "banner": (72, 384, 1456, 360), "mirror": (72, 244, 1000, 500)}
                x, y, w, h = boxes.get(lay, (528, 244, 1000, 500))
                h += 0 if s.get("desk") else 74  # no desk band: the visual takes its room
                slide.shapes.add_picture(str(visuals[s["n"]]), _px(x), _px(y), _px(w), _px(h))
            if not s.get("desk"):
                _text(slide, 72, 850, 1200, 30, s.get("source", ""), 10, TEXT2)
                _notes(slide, s)
                continue
            band = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, _px(72), _px(762), _px(1456), _px(64))
            band.adjustments[0] = 0.18
            band.fill.solid()
            band.fill.fore_color.rgb = _rgb(ACCENT_SOFT)
            band.line.fill.background()
            edge = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, _px(72), _px(762), _px(7), _px(64))
            edge.fill.solid()
            edge.fill.fore_color.rgb = _rgb(ACCENT)
            edge.line.fill.background()
            _text(slide, 98, 784, 220, 24, s.get("desk_label", "For the desk"), 9.5, ACCENT_INK, bold=True, caps=True)
            _text(slide, 320, 774, 1190, 50, s.get("desk", ""), 12.5)
            _text(slide, 72, 850, 1200, 30, s.get("source", ""), 10, TEXT2)
        _notes(slide, s)
    prs.save(out)
    return out


def _visual_lines(v: dict) -> list[str]:
    """The words inside the text-only visuals, so the outline carries everything a template needs."""
    kind = v.get("kind")
    if kind == "primer":
        out = []
        for i, st in enumerate(v["steps"], 1):
            words = "".join(f" {st[k]}" for k in ("formula", "body", "foot") if st.get(k))
            out.append(f"{i}. **{st['title']}**{words}")
            for q in st.get("quotes", []):
                out.append(f"   - {q['label']}: “{q['text']}” ({q['when']}, Claude {q['score']:+.2f})")
        return out
    if kind == "brief":
        return [f"- {ask} → {did}" for ask, did in v["rows"]]
    if kind == "agenda":
        return [f"- {n} {t}: {d}" for n, t, d in v["sections"]]
    if kind == "answers":
        return [f"- {h}: {claim} ({verdict}; {metric})" for h, claim, verdict, metric in v["cards"]]
    if kind == "cards":
        return [f"- **{t}** {d}" for t, d in v["cards"]]
    if kind == "tiles":
        return [f"- **{k}** {d}" for k, d in v["tiles"]]
    if kind == "steps":
        return [f"{i}. **{t}** {d}" for i, (t, d) in enumerate(v["steps"], 1)]
    if kind == "intervals":
        digits = v.get("digits", 3)
        out = [
            f"Axis: {v['axis']}. {v['interval']}.",
            "| Variant | Estimate | Interval | Details |",
            "|---|---:|---|---|",
        ]
        for row in v["rows"]:
            chosen = " (chosen)" if row.get("highlight") and not v.get("zero") else ""
            out.append(
                f"| {row['label']}{chosen} | {row['value']:.{digits}f} | "
                f"{row['low']:.{digits}f} to {row['high']:.{digits}f} | {row.get('details', '')} |"
            )
        return out
    if kind == "table":
        return ["| " + " | ".join(v["columns"]) + " |", "|" + "---|" * len(v["columns"])] + [
            "| " + " | ".join(r) + " |" for r in v["rows"]
        ]
    if kind == "code":
        return ["```text", *v["lines"], "```"]
    if kind == "matrices":
        head = "| expert \\ reader | " + " | ".join(v["labels"]) + " |"
        out = []
        for m in v["items"]:
            out += [f"**{m['title']}**", "", head, "|" + "---|" * (len(v["labels"]) + 1)]
            out += ["| " + " | ".join([v["labels"][i], *map(str, row)]) + " |" for i, row in enumerate(m["m"])]
            out.append("")
        return out
    if kind == "architecture":
        return [
            "- Sources: " + "; ".join(t for t, _ in v["sources"]),
            "- Pipeline: " + " → ".join(f"{t} ({d})" for t, d in v["stages"]),
            "- Store: " + "; ".join(v["store"]),
            "- Ways in: " + "; ".join(f"{t} ({d})" for t, d in v["serve"]),
        ]
    return []


def pdf(images: dict[int, Path], out: Path, story: dict | None = None) -> Path:
    """The rendered slides as one PDF, page for page as they appear on screen (pages are images, not text). Given
    the story, every page gets a bookmark with its slide title, and the file its title and authors."""
    from PIL import Image

    pages = [Image.open(images[n]).convert("RGB") for n in sorted(images)]
    out.parent.mkdir(parents=True, exist_ok=True)
    pages[0].save(out, save_all=True, append_images=pages[1:], resolution=pages[0].width / 13.333, quality=92)
    if story:
        from pypdf import PdfWriter

        doc = PdfWriter(clone_from=out)
        titles = {s["n"]: s["title"] for s in story["slides"]}
        for i, n in enumerate(sorted(images)):
            doc.add_outline_item(f"{n}. {titles.get(n, '')}".rstrip(". "), i)
        cover = story["slides"][0].get("subtitle", "")
        doc.add_metadata({"/Title": "FOMC Tone Lab", "/Author": ", ".join(story["team"]), "/Subject": cover})
        doc.page_mode = "/UseOutlines"
        with open(out, "wb") as f:
            doc.write(f)
    return out


def contact_sheet(images: dict[int, Path], out: Path, cols: int = 4, width: int = 480, gap: int = 16) -> Path:
    from PIL import Image

    h, rows = round(width * STAGE_H / STAGE_W), -(-len(images) // cols)
    sheet = Image.new("RGB", (cols * width + (cols + 1) * gap, rows * h + (rows + 1) * gap), "#e2ded5")
    for k, n in enumerate(sorted(images)):
        thumb = Image.open(images[n]).convert("RGB").resize((width, h), Image.LANCZOS)
        sheet.paste(thumb, (gap + (k % cols) * (width + gap), gap + (k // cols) * (h + gap)))
    sheet.save(out, quality=85, optimize=True)
    return out


def outline(story: dict, out: Path) -> Path:
    name = "the technical appendix" if story.get("deck") == "technical" else "the story"
    lines = [f"# FOMC Tone Lab: {name} ({story['as_of']})", "", f"Team: {', '.join(story['team'])}", ""]
    for s in story["slides"]:
        lines += [f"## {s['n']}. {s['title']}", ""]
        if s.get("kicker"):
            lines.append(f"*{s['kicker']}*")
        if s.get("number"):
            lines.append(f"**{s['number']}**: {s.get('number_label', '')}")
        for c in [c for c in s.get("chips", []) if c]:
            lines.append(f"- {c}")
        lines += _visual_lines(s.get("visual") or {})
        for key in ("subtitle", "desk"):
            if s.get(key):
                lines.append(f"{s.get('desk_label', 'For the desk') + ': ' if key == 'desk' else ''}{s[key]}")
        if s.get("source"):
            lines.append(f"Source: {s['source']}")
        lines += ["", f"> {s.get('notes', '')}", ""]
    out.write_text("\n".join(lines))
    return out


# ------------------------------------------------------------------------ command
def build(
    out_dir: Path, base: str | None = None, workers: int = 4, deck: str = "main", pdf_dir: Path | None = None
) -> dict[str, Path]:
    """One deck as PowerPoint (images and editable), an outline and a contact sheet; deck="technical" writes the
    technical appendix beside the main deck, with -technical in each name. `pdf_dir`: also a PDF there (kept out of
    the repository)."""
    out_dir.mkdir(parents=True, exist_ok=True)
    work = Path(tempfile.mkdtemp(prefix="fomc-slides-"))
    server = None if base else LocalServer().__enter__()
    tag, query = ("-technical", "deck=technical&") if deck == "technical" else ("", "")
    try:
        url = base or server.base
        story = json.loads(urllib.request.urlopen(f"{url}/api/story?deck={deck}", timeout=60).read())
        jobs = []
        for s in story["slides"]:
            n = s["n"]
            jobs.append(
                (
                    f"{url}/?theme=light#/story?{query}s={n}&export=slide",
                    work / f"slide-{n:02d}.png",
                    (STAGE_W, STAGE_H),
                    1.5,
                    ("slide", n),
                )
            )
            if s.get("visual") and s.get("layout") in ("standard", "wide", "mirror", "banner", "compare"):
                w, h = {"wide": (1456, 500), "banner": (1456, 360)}.get(s["layout"], (1000, 500))
                size = (w, h + (0 if s.get("desk") else 74))
                jobs.append(
                    (
                        f"{url}/?theme=light#/story?{query}s={n}&export=visual",
                        work / f"visual-{n:02d}.png",
                        size,
                        2,
                        ("visual", n),
                    )
                )
        images, visuals = {}, {}
        with ThreadPoolExecutor(max_workers=workers) as pool:
            for (kind, n), path in zip([j[4] for j in jobs], pool.map(lambda j: shot(j[0], j[1], j[2], j[3]), jobs)):
                (images if kind == "slide" else visuals)[n] = path
        # Convert the existing vector brand asset for the native PowerPoint cover.
        # All text remains editable; only the original logo is embedded as an image.
        from urllib.parse import quote

        from PIL import Image

        svg = (Path(__file__).resolve().parents[2] / "docs/img/logo.svg").read_text()
        logo = work / "logo.png"
        logo_page = '<html><body style="margin:0;background:#fffefb"><div style="width:84px;height:84px">'
        logo_page += svg.replace('width="64" height="64"', 'width="84" height="84"') + "</div></body></html>"
        shot("data:text/html;charset=utf-8," + quote(logo_page), logo, (STAGE_W, STAGE_H), 2)
        with Image.open(logo) as original:
            original.crop((0, 0, 168, 168)).save(logo)
        return {
            **({"pdf": pdf(images, pdf_dir / f"FOMC-Tone-Lab{tag}.pdf", story)} if pdf_dir else {}),
            "deck": image_deck(story, images, out_dir / f"FOMC-Tone-Lab{tag}.pptx"),
            "editable": editable_deck(story, visuals, out_dir / f"FOMC-Tone-Lab{tag}-editable.pptx", logo=logo),
            "outline": outline(story, out_dir / f"outline{tag}.md"),
            "sheet": contact_sheet(images, out_dir / f"contact-sheet{tag}.jpg"),
            "renders": work,
        }
    finally:
        if server:
            server.__exit__(None, None, None)
